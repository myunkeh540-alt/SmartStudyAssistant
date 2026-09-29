"""
Text extraction helpers for Smart Study Assistant.
"""

import re
import csv
import io

import docx
import openpyxl
from pptx import Presentation
from pypdf import PdfReader

import ai_utils

TEXT_EXTENSIONS = {"txt", "pdf", "docx", "pptx", "csv", "xlsx", "md"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg"}


class TextExtractionError(Exception):
    """Raised when a file's text can't be extracted, with a message safe to show to the user."""


def extract_text_from_pdf(uploaded_file) -> str:
    """Extract plain text from an uploaded PDF file (Streamlit UploadedFile)."""
    reader = PdfReader(uploaded_file)
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages_text).strip()


def extract_text_from_txt(uploaded_file) -> str:
    """Extract plain text from an uploaded TXT file (Streamlit UploadedFile)."""
    return uploaded_file.read().decode("utf-8", errors="ignore").strip()


def extract_text_from_markdown(uploaded_file) -> str:
    """Extract plain text from an uploaded Markdown file (source is already readable)."""
    return extract_text_from_txt(uploaded_file)


def extract_text_from_docx(uploaded_file) -> str:
    """
    Extract paragraph and table text from an uploaded DOCX file as Markdown,
    mapping heading styles to '#'/'##'/... and list styles to '- ' bullets so
    the document's structure survives.
    """
    document = docx.Document(uploaded_file)
    parts = []
    prev_was_list = False

    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "") if p.style else ""

        if style.startswith("Heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            level = min(int(level), 6)
            parts.append(f"{'#' * level} {text}")
            prev_was_list = False
        elif "List" in style:
            parts.append(f"- {text}")
            prev_was_list = True
        else:
            if parts and not prev_was_list:
                parts.append("")
            parts.append(text)
            prev_was_list = False

    for table in document.tables:
        if not table.rows:
            continue
        parts.append("")
        header = [cell.text.strip() for cell in table.rows[0].cells]
        parts.append("| " + " | ".join(header) + " |")
        parts.append("| " + " | ".join("---" for _ in header) + " |")
        for row in table.rows[1:]:
            parts.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")

    return "\n".join(parts).strip()


def extract_text_from_pptx(uploaded_file) -> str:
    """
    Extract text from every slide of an uploaded PPTX file as Markdown: each
    slide's title becomes a heading, other text-frame paragraphs become
    indented bullet points.
    """
    presentation = Presentation(uploaded_file)
    parts = []

    for slide in presentation.slides:
        title_shape = slide.shapes.title
        title_text = title_shape.text.strip() if title_shape and title_shape.has_text_frame else ""
        if title_text:
            parts.append(f"## {title_text}")

        for shape in slide.shapes:
            if shape == title_shape or not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                indent = "  " * (paragraph.level or 0)
                parts.append(f"{indent}- {text}")

    return "\n".join(parts).strip()


def extract_text_from_csv(uploaded_file) -> str:
    """Extract an uploaded CSV file as a Markdown table."""
    decoded = uploaded_file.read().decode("utf-8", errors="ignore")
    rows = [row for row in csv.reader(io.StringIO(decoded)) if any(cell.strip() for cell in row)]
    if not rows:
        return ""

    parts = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join("---" for _ in rows[0]) + " |"]
    for row in rows[1:]:
        parts.append("| " + " | ".join(row) + " |")
    return "\n".join(parts).strip()


def extract_text_from_xlsx(uploaded_file) -> str:
    """Extract every sheet of an uploaded XLSX file as a Markdown table under a heading."""
    workbook = openpyxl.load_workbook(uploaded_file, data_only=True)
    parts = []

    for sheet in workbook.worksheets:
        rows = [
            [str(cell) for cell in row if cell is not None]
            for row in sheet.iter_rows(values_only=True)
            if any(cell is not None for cell in row)
        ]
        if not rows:
            continue
        parts.append(f"### Sheet: {sheet.title}")
        parts.append("| " + " | ".join(rows[0]) + " |")
        parts.append("| " + " | ".join("---" for _ in rows[0]) + " |")
        for row in rows[1:]:
            parts.append("| " + " | ".join(row) + " |")
        parts.append("")

    return "\n".join(parts).strip()


_EXTRACTORS = {
    "txt": extract_text_from_txt,
    "pdf": extract_text_from_pdf,
    "docx": extract_text_from_docx,
    "pptx": extract_text_from_pptx,
    "csv": extract_text_from_csv,
    "xlsx": extract_text_from_xlsx,
    "md": extract_text_from_markdown,
}


_LINE_BREAKS = re.compile(r"[\x0b\x0c  ]")
_JUNK_CHARS = re.compile(r"[\x00-\x08\x0e-\x1f\x7f�​-‍﻿]")


def clean_text(text: str) -> str:
    """Replace in-paragraph line breaks with spaces and drop control characters
    (common in PPTX/DOCX) that would otherwise show up as empty boxes."""
    return _JUNK_CHARS.sub("", _LINE_BREAKS.sub(" ", text))


def extract_text(uploaded_file) -> str:
    """
    Single entry point for the file uploader: routes an uploaded file to the
    right extractor based on its extension, raising TextExtractionError with
    a clear message for unsupported or unreadable files.
    """
    ext = uploaded_file.name.rsplit(".", 1)[-1].lower() if "." in uploaded_file.name else ""

    if ext in IMAGE_EXTENSIONS:
        try:
            return ai_utils.extract_text_from_image(uploaded_file)
        except ai_utils.AIGenerationError as e:
            raise TextExtractionError(str(e)) from e

    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        raise TextExtractionError(f"Unsupported file type: .{ext or 'unknown'}")

    try:
        return clean_text(extractor(uploaded_file))
    except Exception as e:
        raise TextExtractionError(f"Could not read this .{ext} file: {e}") from e
